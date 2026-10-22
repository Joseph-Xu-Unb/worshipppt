import base64
import copy
import json
from pathlib import Path
import tempfile
import unittest

from pptx import Presentation

from src.backend.app import (
    _validate_data,
    generate_worship_powerpoint,
    get_sample_worship_data,
)
from src.backend.content import section_pages, split_lines, validate_section


ROOT = Path(__file__).resolve().parents[1]


class PaginationTests(unittest.TestCase):
    def test_markers_stay_after_full_and_partial_pages(self):
        for marker in ("间奏", "（间奏）", "( 间奏 )", "结束", "（尾声）", "前奏"):
            with self.subTest(marker=marker):
                self.assertEqual(split_lines(["A", "B", "C", "D", marker, "E"]),
                                 [["A", "B", "C", "D", marker], ["E"]])
                self.assertEqual(split_lines(["A", "", "B", marker, "C", "D"]),
                                 [["A", "B", marker], ["C", "D"]])

    def test_markers_do_not_consume_lyric_slots(self):
        self.assertEqual(split_lines(["A", "间奏", "B", "C", "D", "E", "F", "结束"]),
                         [["A", "间奏"], ["B", "C", "D", "E"], ["F", "结束"]])
        self.assertEqual(split_lines(["A", "间奏", "尾声", "B"]),
                         [["A", "间奏", "尾声"], ["B"]])

    def test_orphan_marker_is_rejected(self):
        self.assertIsNotNone(validate_section({"title": "Test", "lines": ["（间奏）", "A"]}, "song"))

    def test_explicit_pages_keep_boundaries(self):
        song = {"title": "Test", "pages": [
            {"lines": ["A", "B"], "footer": "（间奏）"}, {"lines": ["C"]},
            {"lines": ["D", "E"], "footer": "结束"},
        ]}
        self.assertIsNone(validate_section(song, "song"))
        self.assertEqual(section_pages(song), [["A", "B", "（间奏）"], ["C"], ["D", "E", "结束"]])
        song["lines"] = ["different"]
        self.assertIsNotNone(validate_section(song, "song"))

    def test_invalid_pages_are_rejected(self):
        for pages in (None, [], [{}], [{"lines": []}], [{"lines": [" "]}],
                      [{"lines": ["A", "间奏"]}], [{"lines": ["A"], "footer": 3}],
                      [{"lines": ["A"], "footer": "unknown"}]):
            with self.subTest(pages=pages):
                self.assertIsNotNone(validate_section({"title": "Test", "pages": pages}, "song"))

    def test_sample_and_legacy_api_validation(self):
        legacy = json.loads((ROOT / "data/sample_worship_data.json").read_text())
        self.assertIsNone(_validate_data(legacy))
        sample = get_sample_worship_data()
        self.assertIn("pages", sample["hymns"][0])
        self.assertIsNone(_validate_data(sample))
        sample["response_hymn"]["pages"] = [{"lines": ["间奏"]}]
        self.assertIsNotNone(_validate_data(sample))

    def test_actual_pptx_contains_exact_pages_for_both_formats(self):
        base = get_sample_worship_data()
        for explicit in (False, True):
            data = copy.deepcopy(base)
            if explicit:
                song = {"title": "TEST_SONG", "pages": [
                    {"lines": ["TEST_A", "TEST_B"], "footer": "（间奏）"},
                    {"lines": ["TEST_C", "TEST_D"], "footer": "结束"},
                ]}
            else:
                song = {"title": "TEST_SONG", "lines": [
                    "TEST_A", "TEST_B", "（间奏）", "TEST_C", "TEST_D", "结束",
                ]}
            data["hymns"] = [song]
            data["response_hymn"] = song
            self.assertIsNone(_validate_data(data))
            with tempfile.TemporaryDirectory() as directory:
                output = Path(directory) / "test.pptx"
                result = generate_worship_powerpoint(data, "2026-10-25")
                output.write_bytes(base64.b64decode(result["pptx_base64"]))
                self.assertTrue(result["filename"].endswith("2026-10-25.pptx"))
                bodies = []
                for slide in Presentation(output).slides:
                    for shape in slide.shapes:
                        if shape.has_text_frame and (
                            "TEST_A" in shape.text or "TEST_C" in shape.text
                        ):
                            bodies.append(shape.text)
                self.assertEqual(bodies, ["TEST_A\nTEST_B\n（间奏）", "TEST_C\nTEST_D\n结束"] * 2)


if __name__ == "__main__":
    unittest.main()
